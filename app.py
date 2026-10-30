"""Cardio Care AI - single-file Streamlit app.
Run:  pip install -r requirements.txt   then   streamlit run app.py"""
import io, re, json, time
from pathlib import Path
import numpy as np, pandas as pd, streamlit as st

_LOGO = next((p for p in (Path(__file__).parent / n for n in ("logo.png", "logo.jpg", "logo.jpeg")) if p.exists()), Path(__file__).parent / "logo.png")
try:
    from PIL import Image as _Img
    _icon = _Img.open(_LOGO) if _LOGO.exists() else None
except Exception:
    _icon = None
st.set_page_config("Cardio Care AI", _icon, layout="wide", initial_sidebar_state="expanded")

# ======================= REPORT READING & PARAMETER EXTRACTION =======================

import io, re, json
from pathlib import Path

EXT = ["pdf", "png", "jpg", "jpeg", "tif", "tiff", "bmp", "webp", "txt", "md", "docx",
       "xlsx", "xls", "csv", "json", "xml", "html", "htm"]
IMG = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}

# key: (alias regex, min, max, {unit: factor to standard unit})
S = {
    "age": (r"\bage\b", 1, 120, {}),
    "height_cm": (r"\bheight\b", 100, 230, {"in": 2.54, "m": 100}),
    "weight_kg": (r"\bweight\b|\bwt\b", 25, 300, {"lb": 0.4536, "lbs": 0.4536}),
    "hr": (r"heart\s*rate|pulse(?:\s*rate)?|ventricular\s*rate|\bhr\b", 25, 250, {}),
    "hrv": (r"\bhrv\b|rmssd", 3, 300, {}),
    "qtc": (r"\bqtc\b|corrected\s*qt", 280, 700, {}),
    "st_mv": (r"st\s*(?:segment|deviation|level)", -5, 5, {}),
    "spo2": (r"spo2|spo₂|oxygen\s*saturation|o2\s*sat", 50, 100, {}),
    "total_chol": (r"total\s*cholesterol|cholesterol,?\s*total|\btc\b", 50, 600, {"mmol": 38.67}),
    "hdl": (r"(?<!non-)(?<!non )\bhdl(?:[- ]?c)?(?:\s*cholesterol)?", 10, 150, {"mmol": 38.67}),
    "ldl": (r"\bldl(?:[- ]?c)?(?:\s*cholesterol)?", 10, 400, {"mmol": 38.67}),
    "tg": (r"triglycerides?|\btg\b", 20, 2000, {"mmol": 88.57}),
    "lpa": (r"(?<![a-z])lipoprotein\s*\(\s*a\s*\)|(?<![a-z])lp\s*\(\s*a\s*\)|\blpa\b|(?<![a-z])lipoprotein[\s\-]*a\b(?!\s*-?\d)", 0, 500, {}),
    "hba1c": (r"hba1c|glycated\s*h(?:a)?emoglobin|\ba1c\b", 3, 18, {}),
    "glucose": (r"(?:fasting\s*)?(?:blood\s*)?glucose|\bfbs\b|\bfbg\b", 40, 600, {"mmol": 18.016}),
    "creatinine": (r"creatinine", 0.2, 15, {"µmol": 0.0113, "umol": 0.0113}),
    "lvef": (r"ejection\s*fraction|\blvef\b|\bef\b", 5, 85, {}),
    "fai": (r"fat\s*attenuation(?:\s*index)?|\bfai\b", -140, 0, {}),
    "cac": (r"calcium\s*score|\bcac\b|agatston", 0, 10000, {}),
    "pr": (r"\bpr\s*interval", 80, 400, {}),
    "qrs": (r"\bqrs(?:\s*(?:duration|interval|width))?", 40, 250, {}),
    "troponin": (r"hs[- ]?c?tn[it]\b|troponin(?:\s*[- ]?\s*[it]\b)?", 0, 100000, {"ng/ml": 1000, "ug/l": 1000, "µg/l": 1000}),
    "ckmb": (r"ck[- ]?mb", 0, 1000, {}),
    "bnp": (r"nt[- ]?pro[- ]?bnp|\bbnp\b", 0, 100000, {}),
    "hscrp": (r"hs[- ]?crp|c[- ]?reactive\s*protein|\bcrp\b", 0, 500, {"mg/dl": 10}),
    "egfr": (r"\begfr\b|estimated\s*gfr", 5, 200, {}),
    "hemoglobin": (r"(?<!glycated\s)(?<!glycosylated\s)h(?:a)?emoglobin(?!\s*a1c)|\bhb\b(?!a1c)", 3, 25, {"g/l": 0.1}),
    "apob": (r"apo\s*lipoprotein\s*b|\bapo[- ]?b\b", 20, 300, {}),
    "apoa1": (r"apo\s*lipoprotein\s*a-?1|\bapo[- ]?a-?1\b", 30, 300, {}),
    "non_hdl": (r"non[- ]?hdl(?:[- ]?c)?(?:\s*cholesterol)?", 20, 500, {"mmol": 38.67}),
}
FLAGS = {
    "diabetes": r"diabet|\bt2dm\b|\bdm\b",
    "htn": r"hypertensi|\bhtn\b",
    "fam_hx": r"family\s*history\s*(?:of\s*)?(?:cad|coronary|heart|premature|mi)\w*",
    "prior_mi": r"myocardial\s*infarction|\bmi\b|heart\s*attack|\bn?stemi\b",
    "prior_pci": r"\bpci\b|angioplasty|\bstent",
    "prior_cabg": r"\bcabg\b|bypass\s*(?:graft|surgery)",
    "stroke": r"\bstroke\b(?!\s*volume)|\bcva\b|\btia\b",
    "ckd": r"chronic\s*kidney|\bckd\b",
    "dyslip": r"dyslipid|hyperlipid|hypercholesterol",
    "afib": r"atrial\s*fibrillation|\bafib\b|\baf\b",
    "lvh": r"left\s*ventricular\s*hypertrophy|\blvh\b",
    "wma": r"wall\s*motion\s*abnormal|hypokines|akines|dyskines",
    "lbbb": r"\blbbb\b|bundle\s*branch\s*block",
    "chest_pain": r"chest\s*pain|angina",
    "dyspnea": r"dyspn|shortness\s*of\s*breath|breathless",
    "palp": r"palpitation",
    "syncope": r"syncope|fainting",
    "abd_obesity": r"abdominal\s*obesity|central\s*obesity",
    "fatigue": r"fatigue|tiredness",
}
NUM = r"(-?\d+(?:[.,]\d+)?)"


def _ocr(img):
    import pytesseract
    return pytesseract.image_to_string(img)


def read_text(name, data):
    """Return (text, notes). Never raises: problems go into notes."""
    ext, notes = Path(name).suffix.lower(), []
    try:
        if ext == ".pdf":
            import pdfplumber
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                text = "\n".join((p.extract_text() or "") for p in pdf.pages)
            if len(text.strip()) < 40:  # scanned PDF -> OCR
                import pypdfium2
                pdf = pypdfium2.PdfDocument(data)
                text = "\n".join(_ocr(pdf[i].render(scale=2.5).to_pil()) for i in range(len(pdf)))
                notes.append("Scanned PDF read with OCR")
            return text, notes
        if ext in IMG:
            from PIL import Image
            notes.append("Image read with OCR")
            return _ocr(Image.open(io.BytesIO(data))), notes
        if ext == ".docx":
            import docx
            d = docx.Document(io.BytesIO(data))
            lines = [p.text for p in d.paragraphs]
            for t in d.tables:
                lines += ["  ".join(c.text.strip() for c in r.cells) for r in t.rows]
            return "\n".join(lines), notes
        if ext in (".xlsx", ".xls", ".csv"):
            import pandas as pd
            sheets = ({"": pd.read_csv(io.BytesIO(data), header=None, dtype=str)} if ext == ".csv"
                      else pd.read_excel(io.BytesIO(data), sheet_name=None, header=None, dtype=str))
            return "\n".join("  ".join(str(c) for c in row if str(c) != "nan")
                             for df in sheets.values() for row in df.values.tolist()), notes
        if ext == ".json":
            def flat(o, p=""):
                if isinstance(o, dict):
                    return [x for k, v in o.items() for x in flat(v, k)]
                if isinstance(o, list):
                    return [x for v in o for x in flat(v, p)]
                return [f"{p}: {o}"]
            return "\n".join(flat(json.loads(data.decode("utf-8", "ignore")))), notes
        text = data.decode("utf-8", "ignore")
        if ext in (".xml", ".html", ".htm"):
            text = re.sub(r"<[^>]+>", "\n", text)
        return text, notes
    except Exception as e:  # missing library, corrupt file, no tesseract...
        return "", notes + [f"Could not read file: {e}"]


def _flag(text, terms):
    found = None
    for m in re.finditer(terms, text, re.I):
        pre = text[max(0, m.start() - 35):m.start()].split("\n")[-1]
        post = text[m.end():m.end() + 14]
        neg = (re.search(r"(?:\bno|denies|without|negative for|nil|absent|\bnot)\b[^.]*$", pre, re.I)
               or re.match(r"\s*[:\-]?\s*(?:no|none|nil|negative|absent)\b", post, re.I))
        if not neg:
            return True
        found = False
    return found


_TITLE = r"(?:mr|mrs|ms|miss|dr|master|baby|shri|smt|sri|prof)"
_NOT_NAME = {"details", "detail", "information", "info", "report", "history", "id", "number", "no", "name", "type", "data", "summary", "sample", "profile",
             "date", "age", "sex", "gender", "test", "lab", "laboratory", "hospital", "clinic", "diagnostics", "centre", "center", "pathology", "male", "female",
             "years", "year", "yrs", "unknown", "na", "nil", "none", "of", "and", "the"}


def _clean_name(s):
    """Cut a captured line down to just the name; None if what is left does not look like a name."""
    s = re.split(r"\s{2,}|\t|\||,|;|\(|/|\d|\s(?:age|sex|gender|dob|d\.o\.b|id|uhid|mrn|reg|ref|date|lab|sample|ward|bed|opd|ipd|phone|mobile|tel|address|doctor|referred|collected|reported)\b",
                 s.strip(), flags=re.I)[0].strip(" .-:–=")
    words = re.findall(r"[A-Za-z][A-Za-z.'\-]*", s)
    if not 1 <= len(words) <= 5 or not 2 <= len(s) <= 45 or not re.fullmatch(r"[A-Za-z][A-Za-z .'\-]*", s):
        return None
    return None if any(w.lower().strip(".") in _NOT_NAME for w in words) else s


def _find_name(text):
    """(name, matched text, confidence) from labels such as 'Patient Name: Mr. Ravi Kumar', 'Name  RAVI KUMAR', a name on the
    line after the label, or a title like 'Mrs. Lakshmi Devi'. None if no name is found."""
    strong = r"(?:patient(?:'s)?\s*name|name\s*of\s*(?:the\s*)?patient|pt\.?\s*name)"
    weak = (r"(?:patient|(?<!doctor )(?<!physician )(?<!consultant )(?<!test )(?<!lab )(?<!ref )(?<!referred by )(?<!sample )(?<!hospital )\bname)")
    pats = ((strong + r"[ \t]*[:\-–|=]*[ \t]*([^\n]{2,70})", "high"),             # label and value on one line, separator optional
            (weak + r"[ \t]*(?:[:\-–|=]+|\t| {2,})[ \t]*([^\n]{2,70})", "high"),  # weaker label needs a separator
            (strong + r"[ \t]*[:\-–]?[ \t]*\n\s*([^\n]{2,70})", "high"))           # value on the next line
    for pat, conf in pats:
        for m in re.finditer(pat, text, re.I):
            t = re.match(rf"\s*({_TITLE}\.?\s+)", m[1], re.I)
            nm = _clean_name(m[1][t.end():] if t else m[1])
            if nm:
                return ((t[1] if t else "") + nm).strip(), m[0], conf
    # more layouts: single-space PDF tables ('Name Ravi Kumar Age 52'), 'Name/Age/Sex: Ravi Kumar/52/M', 'Client Name'
    for m in re.finditer(r"\b(?:name|patient|client|customer)(?:[ \t]*[/&][ \t]*(?:age|sex|gender|id|uhid|mrn))*[ \t]*[:\-–|=]?[ \t]+([A-Za-z][A-Za-z.'\-]*(?:[ \t][A-Za-z][A-Za-z.'\-]*){0,4})", text, re.I):
        before = text[max(0, m.start() - 16):m.start()].lower()
        if not m[1][0].isupper() or re.search(r"(?:doctor|physician|consultant|test|lab|ref|referred by|sample|hospital|clinic|father|mother|husband|spouse|guardian|dr\.?|insurer|company)[ \t]$", before):
            continue
        t = re.match(rf"\s*({_TITLE}\.?\s+)", m[1], re.I)
        nm = _clean_name(m[1][t.end():] if t else m[1])
        if nm:
            return ((t[1] if t else "") + nm).strip(), m[0], "medium"
    # a line that starts with a name followed by age and sex, e.g. 'RAVI KUMAR 52 Y / M' or 'Ravi Kumar (52 Years, Male)'
    if m := re.search(r"^[ \t]*((?:" + _TITLE + r"\.?[ \t]+)?[A-Z][A-Za-z.'\-]+(?:[ \t][A-Z][A-Za-z.'\-]+){1,3})[ \t,(]+\d{1,3}[ \t]*(?i:y|yrs?|years?)?[ \t]*[/,\-]?[ \t]*(?i:m|f|male|female)\b", text, re.M):
        nm = _clean_name(m[1])
        if nm:
            return nm, m[0], "medium"
    if m := re.search(r"\b(?:Mr|MR|Mrs|MRS|Ms|MS|Miss|MISS|Shri|SHRI|Smt|SMT|Master|MASTER)\.?[ \t]+([A-Z][A-Za-z]+(?:[ \t][A-Z][A-Za-z]+){0,2})", text):
        return m[0], m[0], "medium"
    return None


def parse(text):
    """Return ({key: value}, hits). hits = [{key,value,raw,conf}] for the audit table."""
    out, hits = {}, []

    def put(k, v, raw, conf="high"):
        if k not in out:
            out[k] = v
            hits.append({"key": k, "value": v, "raw": raw.strip()[:90], "conf": conf})

    if m := re.search(r"(?:blood\s*pressure|\bbp\b)[^\d\n]{0,15}(\d{2,3})\s*/\s*(\d{2,3})", text, re.I):
        s, d = int(m[1]), int(m[2])
        if 60 <= s <= 260 and 30 <= d <= 150 and s > d:
            put("sbp", s, m[0]); put("dbp", d, m[0])
    for k, (alias, lo, hi, conv) in S.items():
        for m in re.finditer(alias, text, re.I):
            tail = text[m.end():].split("\n", 1)[0][:80]
            n = re.search(NUM, tail)
            if not n:
                continue
            v = float(n[1].replace(",", "."))
            unit = tail[n.end():n.end() + 14].lower().lstrip()
            if k == "height_cm" and (f := re.match(r"\s*(?:ft|')\s*(\d{1,2})", unit)):
                v = v * 30.48 + int(f[1]) * 2.54
            elif k == "height_cm" and unit.startswith(("ft", "'")):
                v *= 30.48
            elif k == "hba1c" and unit.startswith("mmol/mol"):
                v = v / 10.929 + 2.15
            else:
                for u, f in conv.items():
                    if unit.startswith(u):
                        v *= f; break
            v = round(v, 2)
            if lo <= v <= hi:
                put(k, int(v) if v == int(v) and k in ("age", "hr", "total_chol", "hdl", "ldl", "tg", "cac", "qtc", "pr", "qrs", "stenosis") else v,
                    m.group(0) + tail[:n.end()], "high" if unit[:1].isalpha() or unit[:1] == "%" else "medium")
                break
    sten = [int(x) for x in re.findall(r"(\d{1,3})\s*%\s*(?:\w+\s){0,2}(?:stenosis|occlusion|blockage)", text, re.I)
            + re.findall(r"(?:stenosis|occlusion)\D{0,15}(\d{1,3})\s*%", text, re.I)]
    if sten and max(sten) <= 100:
        put("stenosis", max(sten), "stenosis " + str(max(sten)) + "%")
    if m := re.search(r"(sinus\s*(?:rhythm|tachycardia|bradycardia)|atrial\s*(?:fibrillation|flutter))", text, re.I):
        put("rhythm", m[1].strip().capitalize(), m[0])
    if (m := re.search(r"\b(?:sex|gender)\b\s*[:\-/]?\s*(male|female|m|f)\b", text, re.I)) or \
       (m := re.search(r"\b(male|female)\b", text, re.I)):
        put("sex", "Male" if m[1].lower().startswith("m") else "Female", m[0])
    if m := re.search(r"age\s*/\s*(?:sex|gender)\s*[:\-]?\s*(\d{1,3})\s*(?:y\w*)?\s*[/,\-]?\s*(male|female|m|f)\b", text, re.I):
        put("age", int(m[1]), m[0]); put("sex", "Male" if m[2].lower().startswith("m") else "Female", m[0])
    if "age" not in out:
        m = re.search(r"\b(\d{1,3})\s*(?:y|yrs?|years?)\b(?!\s*ago)", text, re.I)
        if m and 1 <= int(m[1]) <= 120:
            put("age", int(m[1]), m[0], "medium")
        elif m := re.search(r"(?:dob|d\.o\.b|date\s*of\s*birth|birth\s*date)\D{0,12}(\d{1,2})[-/. ](\d{1,2})[-/. ](\d{4})", text, re.I):
            import datetime as dt
            t, (d, mo, y) = dt.date.today(), (int(m[1]), int(m[2]), int(m[3]))
            a = t.year - y - ((t.month, t.day) < (mo, d))
            if 0 < a < 120:
                put("age", a, m[0], "medium")
    if nm := _find_name(text):
        put("name", nm[0], nm[1], nm[2])
    if m := re.search(r"(never\s*smok\w*|non[- ]?smoker|former\s*smoker|ex[- ]?smoker|current\s*smoker|smok\w*\s*(?:status)?\s*[:\-]?\s*(?:never|former|current|yes|no))", text, re.I):
        t = m[0].lower()
        put("smoking", "Never" if re.search(r"never|non|:\s*no|-\s*no", t) else
            "Former" if re.search(r"former|ex", t) else "Current", m[0])
    for k, terms in FLAGS.items():
        v = _flag(text, terms)
        if v is not None:
            put(k, v, k, "medium")
    return out, hits


def extract(name, data):
    text, notes = read_text(name, data)
    params, hits = parse(text)
    if not text.strip():
        notes.append("No readable text found")
    elif not params:
        notes.append("Text was read but no known cardiovascular parameters were recognised")
    return {"file": name, "params": params, "hits": hits, "notes": notes, "chars": len(text), "text": text[:3000]}


# ======================= TRAINED CAD MODEL (Z-Alizadeh Sani dataset) =======================
# Every parameter that exists in the dataset is scored by a model trained on it. Parameters that are not in the
# dataset (troponin, calcium score, QTc, HRV ...) stay rule-based and are blended in by risk() below.
from functools import lru_cache

DATA_FILE = "extention_of_Z-Alizadeh_sani_dataset.xlsx"
# app key -> (dataset column, display name)
MODEL_INPUTS = {
    "age": ("Age", "Age"), "sex": ("Sex", "Sex"), "bmi": ("BMI", "BMI"), "diabetes": ("DM", "Diabetes"),
    "htn": ("HTN", "Hypertension"), "smoking": ("Current Smoker", "Current smoking"), "fam_hx": ("FH", "Family history of CAD"),
    "chest_pain": ("Typical Chest Pain", "Chest pain / angina"), "dyspnea": ("Dyspnea", "Breathlessness"),
    "sbp": ("BP", "Systolic BP"), "hr": ("PR", "Heart rate"), "ldl": ("LDL", "LDL cholesterol"), "hdl": ("HDL", "HDL cholesterol"),
    "tg": ("TG", "Triglycerides"), "hemoglobin": ("HB", "Hemoglobin"), "lvef": ("EF-TTE", "LV ejection fraction"),
}
YES_NO_COLS = {"DM", "HTN", "Current Smoker", "FH", "Typical Chest Pain", "Dyspnea"}
MEASURED = ["age", "sex", "bmi", "smoking", "sbp", "hr", "ldl", "hdl", "tg", "hemoglobin", "lvef"]  # yes/no findings are not counted
MIN_INPUTS = 5  # below this many measured inputs the trained model is skipped and the rule-based score is used


def _model_row(d):
    """Map the app's values to the model's inputs. Missing numbers stay NaN (the model fills them with the training
    median); yes/no findings that were not reported count as absent."""
    row = {}
    for k, (col, _) in MODEL_INPUTS.items():
        v = d.get(k)
        if k == "sex": x = np.nan if v is None else 1.0 if v == "Male" else 0.0
        elif k == "smoking": x = np.nan if v is None else 1.0 if v == "Current" else 0.0
        elif k == "diabetes": x = 1.0 if d.get("diabetes") or (d.get("hba1c") or 0) >= 6.5 else 0.0
        elif k == "htn": x = 1.0 if d.get("htn") or (d.get("sbp") or 0) >= 140 or (d.get("dbp") or 0) >= 90 else 0.0
        elif k in ("fam_hx", "chest_pain", "dyspnea"): x = 1.0 if d.get(k) else 0.0
        else: x = np.nan if v is None else float(v)
        row[col] = x
    return row


@st.cache_resource(show_spinner="Training the CAD risk model on the Z-Alizadeh Sani dataset…")
def train_cad_model():
    """Cross-validates five models plus a soft-voting ensemble of the top two on the dataset, keeps the one with the
    highest mean ROC-AUC, refits it on the training split and checks it once on a held-out test split."""
    try:
        path = next((p for p in (Path(__file__).parent / DATA_FILE, Path.cwd() / DATA_FILE) if p.exists()), None)
        if path is None:
            return {"ok": False, "error": f"{DATA_FILE} not found next to app.py"}
        from sklearn.base import clone
        from sklearn.ensemble import GradientBoostingClassifier, RandomForestClassifier, VotingClassifier
        from sklearn.impute import SimpleImputer
        from sklearn.linear_model import LogisticRegression
        from sklearn.metrics import roc_auc_score
        from sklearn.model_selection import StratifiedKFold, cross_val_score, train_test_split
        from sklearn.neighbors import KNeighborsClassifier
        from sklearn.pipeline import Pipeline
        from sklearn.preprocessing import StandardScaler
        from sklearn.svm import SVC

        z = pd.read_excel(path)
        cols = [c for c, _ in MODEL_INPUTS.values()]
        X = z[cols].copy()
        for c in YES_NO_COLS:
            X[c] = z[c].map(lambda t: 1.0 if str(t).strip().upper() in ("Y", "YES", "1", "1.0") else 0.0)
        X["Sex"] = z["Sex"].map(lambda t: 1.0 if str(t).strip().lower().startswith("m") else 0.0)
        X = X.astype(float)
        y = (z["Cath"].astype(str).str.strip().str.upper() == "CAD").astype(int)
        keep = X.notna().all(axis=1)
        X, y = X[keep], y[keep]
        X_tr, X_te, y_tr, y_te = train_test_split(X, y, test_size=0.15, stratify=y, random_state=42)

        pipe = lambda clf: Pipeline([("imp", SimpleImputer(strategy="median")), ("sc", StandardScaler()), ("clf", clf)])
        cand = {
            "Logistic Regression": pipe(LogisticRegression(max_iter=5000, class_weight="balanced", C=0.5)),
            "Random Forest": pipe(RandomForestClassifier(n_estimators=300, max_depth=5, class_weight="balanced", random_state=42)),
            "Gradient Boosting": pipe(GradientBoostingClassifier(n_estimators=150, max_depth=3, random_state=42)),
            "Support Vector Machine (RBF)": pipe(SVC(probability=True, class_weight="balanced", C=1.0, random_state=42)),
            "K-Nearest Neighbors": pipe(KNeighborsClassifier(n_neighbors=7)),
        }
        skf = StratifiedKFold(n_splits=5, shuffle=True, random_state=42)
        scores = {n: cross_val_score(clone(m), X_tr, y_tr, cv=skf, scoring="roc_auc") for n, m in cand.items()}
        top2 = [n for n, _ in sorted(scores.items(), key=lambda kv: np.mean(kv[1]), reverse=True)[:2]]
        ens = f"Ensemble ({top2[0]} + {top2[1]})"
        cand[ens] = VotingClassifier([(n, clone(cand[n])) for n in top2], voting="soft")
        scores[ens] = cross_val_score(clone(cand[ens]), X_tr, y_tr, cv=skf, scoring="roc_auc")

        best = max(scores, key=lambda n: np.mean(scores[n]))
        model = clone(cand[best]).fit(X_tr, y_tr)
        return {
            "ok": True, "model": model, "best": best, "cols": cols, "median": X_tr.median(), "n": int(len(X)),
            "cv": sorted(((n, float(np.mean(s)), float(np.std(s))) for n, s in scores.items()), key=lambda r: -r[1]),
            "test_auc": float(roc_auc_score(y_te, model.predict_proba(X_te)[:, 1])),
        }
    except Exception as e:
        return {"ok": False, "error": str(e)}


@lru_cache(maxsize=4096)
def _score_cached(key):
    """Probability of CAD for one patient plus how much each reported factor pushed it up or down (each factor is
    swapped for the training-set typical value, or for 'absent' if it is a yes/no finding, in one batched call)."""
    M = train_cad_model()
    cols = M["cols"]
    base = pd.Series([np.nan if x is None else x for x in key], index=cols, dtype=float)
    rows, names = [base], []
    for c, x in zip(cols, key):
        if x is None or (c in YES_NO_COLS and x == 0):
            continue
        r = base.copy()
        r[c] = 0.0 if c in YES_NO_COLS else M["median"][c]
        rows.append(r); names.append(c)
    P = M["model"].predict_proba(pd.DataFrame(rows, columns=cols))[:, 1]
    return float(P[0]), {n: float(P[0] - p) for n, p in zip(names, P[1:])}


def cad_model_predict(d):
    """{'prob', 'present', 'drivers'} from the trained model, or None if it could not be trained."""
    M = train_cad_model()
    if not M["ok"]:
        return None
    row = _model_row(d)
    present = sum(not np.isnan(row[MODEL_INPUTS[k][0]]) for k in MEASURED)
    if present < MIN_INPUTS:
        return {"prob": None, "present": present}
    prob, delta = _score_cached(tuple(None if np.isnan(row[c]) else row[c] for c in M["cols"]))
    names = {col: name for col, name in MODEL_INPUTS.values()}
    return {"prob": prob, "present": present, "drivers": {names[c]: x for c, x in delta.items()}}


# ======================= PARAMETERS, SCENARIOS & RISK MODEL =======================
# key, label, unit, min, max, step, normal (low, high), group
PARAMS = [
 ("hr", "Heart rate", "bpm", 25, 250, 1, (60, 100), "Heart & ECG"), ("sbp", "Systolic BP", "mmHg", 60, 260, 1, (90, 119), "Heart & ECG"),
 ("dbp", "Diastolic BP", "mmHg", 30, 150, 1, (60, 79), "Heart & ECG"), ("hrv", "HRV (RMSSD)", "ms", 3, 300, 1, (30, None), "Heart & ECG"),
 ("qtc", "QTc interval", "ms", 280, 700, 1, (350, 450), "Heart & ECG"), ("pr", "PR interval", "ms", 80, 400, 1, (120, 200), "Heart & ECG"),
 ("qrs", "QRS duration", "ms", 40, 250, 1, (70, 110), "Heart & ECG"), ("st_mv", "ST deviation", "mV", -5, 5, 0.01, (-0.05, 0.05), "Heart & ECG"),
 ("spo2", "SpO₂", "%", 50, 100, 1, (95, 100), "Heart & ECG"),
 ("total_chol", "Total cholesterol", "mg/dL", 50, 600, 1, (None, 200), "Lipid profile"), ("ldl", "LDL cholesterol", "mg/dL", 10, 400, 1, (None, 100), "Lipid profile"),
 ("hdl", "HDL cholesterol", "mg/dL", 10, 150, 1, (40, None), "Lipid profile"), ("tg", "Triglycerides", "mg/dL", 20, 2000, 1, (None, 150), "Lipid profile"),
 ("non_hdl", "Non-HDL cholesterol", "mg/dL", 20, 500, 1, (None, 130), "Lipid profile"), ("apob", "Apolipoprotein B", "mg/dL", 20, 300, 1, (None, 90), "Lipid profile"),
 ("apoa1", "Apolipoprotein A1", "mg/dL", 30, 300, 1, (120, None), "Lipid profile"), ("lpa", "Lipoprotein(a)", "mg/dL", 0, 500, 0.1, (None, 30), "Lipid profile"),
 ("troponin", "Troponin (hs)", "ng/L", 0, 100000, 0.1, (None, 14), "Cardiac biomarkers"), ("ckmb", "CK-MB", "ng/mL", 0, 1000, 0.1, (None, 5), "Cardiac biomarkers"),
 ("bnp", "BNP / NT-proBNP", "pg/mL", 0, 100000, 1, (None, 100), "Cardiac biomarkers"), ("hscrp", "hs-CRP", "mg/L", 0, 500, 0.1, (None, 3), "Cardiac biomarkers"),
 ("hba1c", "HbA1c", "%", 3, 18, 0.1, (4, 5.6), "Metabolic & renal"), ("glucose", "Glucose", "mg/dL", 40, 600, 1, (70, 99), "Metabolic & renal"),
 ("creatinine", "Creatinine", "mg/dL", 0.2, 15, 0.01, (0.6, 1.3), "Metabolic & renal"), ("egfr", "eGFR", "mL/min", 5, 200, 1, (60, None), "Metabolic & renal"),
 ("hemoglobin", "Hemoglobin", "g/dL", 3, 25, 0.1, (12, 17.5), "Metabolic & renal"),
 ("lvef", "LV ejection fraction", "%", 5, 85, 1, (50, 75), "Imaging & structure"), ("cac", "Calcium score (CAC)", "AU", 0, 10000, 1, (None, 0), "Imaging & structure"),
 ("fai", "Fat attenuation index", "HU", -140, 0, 0.1, (None, -70.1), "Imaging & structure"), ("stenosis", "Max coronary stenosis", "%", 0, 100, 1, (None, 49), "Imaging & structure"),
]
LAB = {p[0]: (p[1], p[2]) for p in PARAMS}
REF = {p[0]: p[6] for p in PARAMS}
GROUPS = list(dict.fromkeys(p[7] for p in PARAMS))
FLAGLAB = {"diabetes": "Diabetes", "htn": "Hypertension", "dyslip": "Dyslipidemia", "fam_hx": "Family history of CAD", "prior_mi": "Prior heart attack (MI)",
           "prior_pci": "Prior PCI / stent", "prior_cabg": "Prior bypass (CABG)", "stroke": "Stroke / TIA", "ckd": "Chronic kidney disease", "afib": "Atrial fibrillation",
           "lvh": "LV hypertrophy", "wma": "Wall-motion abnormality", "lbbb": "Bundle branch block", "chest_pain": "Chest pain / angina",
           "dyspnea": "Breathlessness", "palp": "Palpitations", "syncope": "Syncope / fainting", "abd_obesity": "Abdominal obesity", "fatigue": "Fatigue"}
KEYS = ["name", "age", "sex", "height_cm", "weight_kg", "smoking", "activity", "diet", "alcohol", "psych_stress", "rhythm"] + [p[0] for p in PARAMS] + list(FLAGLAB)

HEALTHY = dict(name="Demo patient", age=45, sex="Male", height_cm=172, weight_kg=68, smoking="Never", activity="High", diet="High", alcohol="Light", psych_stress="Low",
               qtc=421, st_mv=0.0, spo2=98, pr=160, qrs=90, total_chol=180, hdl=55, tg=120, lpa=12.9, hba1c=5.2, lvef=60, fai=-82, cac=0, troponin=4,
               bnp=40, hscrp=1, egfr=95, rhythm="Sinus rhythm")
CAD = {**HEALTHY, "age": 58, "smoking": "Former", "activity": "Low", "diet": "Moderate", "qtc": 440, "total_chol": 235, "hdl": 38, "tg": 180, "lpa": 40,
       "fai": -68, "cac": 0, "htn": True, "fam_hx": True, "hscrp": 4, "psych_stress": "Moderate"}
# (group, name, description, hr, sbp, dbp, hrv, overrides)
SCENARIOS = [
 ("Healthy", "Healthy — Baseline", "Normal resting state.", 68, 115, 75, 70, HEALTHY),
 ("Healthy", "Healthy — Post-Exercise", "Recovery after moderate exercise.", 110, 148, 86, 25, HEALTHY),
 ("CAD", "CAD — Borderline Hypertensive", "Hypertension with family history.", 82, 142, 91, 38, CAD),
 ("CAD", "CAD — High Stress / Low HRV", "High sympathetic tone, low HRV.", 96, 135, 88, 18, CAD),
 ("CAD", "CAD — Cardiac Concern", "Prolonged QTc and ST changes.", 88, 158, 98, 14, {**CAD, "qtc": 480, "st_mv": -0.15, "chest_pain": True}),
 ("CAD", "CAD — Smoker + Sedentary", "Current smoker, sedentary.", 92, 146, 92, 22, {**CAD, "smoking": "Current"}),
 ("CAD", "CAD — Diabetic + Hypertensive", "Diabetes with Stage 2 hypertension.", 84, 154, 94, 28, {**CAD, "diabetes": True, "hba1c": 7.8}),
 ("CAD", "CAD — Post-MI Recovery", "Post-infarction monitoring.", 72, 124, 80, 45, {**CAD, "prior_mi": True, "prior_pci": True}),
 ("Other CVD", "CVD — Heart Failure (HFrEF)", "Reduced ejection fraction (30%).", 104, 104, 68, 16, {**CAD, "lvef": 30, "bnp": 900, "dyspnea": True}),
 ("Other CVD", "CVD — Atrial Fibrillation", "Irregular rhythm, raised stroke risk.", 118, 132, 84, 12, {**CAD, "afib": True, "rhythm": "Atrial fibrillation"}),
 ("Other CVD", "CVD — Hemorrhagic Stroke", "Severe hypertension, stroke.", 102, 192, 112, 18, {**CAD, "stroke": True}),
 ("Other CVD", "CVD — Stroke (Cerebrovascular)", "Stroke or TIA documented; elevated vascular risk.", 88, 168, 96, 30, {**CAD, "stroke": True}),
]
CORE = ["hr", "sbp", "dbp", "qtc", "total_chol", "hdl", "tg", "age", "sex", "smoking", "lvef", "hba1c"]
MISSING = {k: None for k in KEYS}

def scenario_values(name):
    g, n, d, hr, s, dia, hrv, o = next(x for x in SCENARIOS if x[1] == name)
    return {**MISSING, **o, "hr": hr, "sbp": s, "dbp": dia, "hrv": hrv}

def bp_class(s, d):
    if s is None or d is None: return "—", "#64748b"
    if s >= 180 or d >= 120: return "Hypertensive crisis", "#dc2626"
    if s >= 140 or d >= 90: return "Stage 2 hypertension", "#dc2626"
    if s >= 130 or d >= 80: return "Stage 1 hypertension", "#d97706"
    if s >= 120: return "Elevated", "#d97706"
    return "Normal", "#16a34a"

def derive(v):
    """Adds BMI, non-HDL, LDL (Friedewald), ApoB/ApoA1 estimates, stress index. Reported values always win."""
    d, est, ok = dict(v), set(), lambda *a: all(v.get(x) is not None for x in a)
    if ok("height_cm", "weight_kg"): d["bmi"] = round(v["weight_kg"] / (v["height_cm"] / 100) ** 2, 1)
    if v.get("non_hdl") is None and ok("total_chol", "hdl"): d["non_hdl"] = v["total_chol"] - v["hdl"]; est.add("non_hdl")
    if v.get("ldl") is None and ok("total_chol", "hdl", "tg"): d["ldl"] = round(v["total_chol"] - v["hdl"] - v["tg"] / 5); est.add("ldl")
    if v.get("apob") is None and d.get("non_hdl") is not None: d["apob"] = round(0.65 * d["non_hdl"] + 6.3, 1); est.add("apob")
    if v.get("apoa1") is None and ok("hdl"): d["apoa1"] = round(v["hdl"] * 2.0); est.add("apoa1")
    if d.get("apob") and d.get("apoa1"): d["apob_ratio"] = round(d["apob"] / d["apoa1"], 2)
    if ok("hrv", "hr"): d["stress"] = round(min(100, max(0, 100 - 1.15 * v["hrv"] + max(0, v["hr"] - 70) * 0.3)))
    d["_est"] = est
    return d

def band(x, steps):
    return next((p for t, p in steps if x >= t), 0)

def rule_risk(v):
    """(score 0-100, tier, {factor: points}, completeness %). Transparent illustrative rules - NOT a validated clinical model.
    Used for scenario detection and as the fallback when the trained model is unavailable."""
    d, c = derive(v), {}
    g = lambda k: d.get(k)
    if g("sbp") is not None: c["Blood pressure"] = max(band(g("sbp"), [(180, 18), (140, 12), (130, 6), (120, 3)]), band(g("dbp") or 0, [(120, 18), (90, 12), (80, 6)]))
    if g("hrv") is not None: c["Low HRV"] = 10 if g("hrv") < 20 else 7 if g("hrv") < 30 else 4 if g("hrv") < 42 else 0
    if g("stress") is not None: c["Physiological stress"] = round(g("stress") / 10)
    if g("qtc") is not None: c["QTc interval"] = 10 if g("qtc") > 500 else 5 if g("qtc") > (460 if g("sex") == "Female" else 450) else 1
    if g("qrs") is not None: c["QRS duration"] = 3 if g("qrs") >= 120 else 0
    if g("pr") is not None: c["PR interval"] = 2 if g("pr") > 200 else 0
    if g("st_mv") is not None: c["ST deviation"] = 5 if abs(g("st_mv")) >= 0.1 else 0
    if g("apob") is not None: c["ApoB"] = band(g("apob"), [(130, 14), (100, 9), (80, 4)])
    if g("ldl") is not None: c["LDL cholesterol"] = band(g("ldl"), [(190, 8), (160, 6), (130, 3)])
    if g("total_chol") is not None: c["Total cholesterol"] = band(g("total_chol"), [(240, 6), (200, 3)])
    if g("lpa") is not None: c["Lipoprotein(a)"] = band(g("lpa"), [(50, 4)])
    if g("smoking"): c["Smoking"] = {"Never": 0, "Former": 4, "Current": 10}[g("smoking")]
    if g("age") is not None: c["Age"] = band(g("age"), [(65, 8), (55, 5), (45, 2)])
    if g("diabetes") or (g("hba1c") or 0) >= 6.5: c["Diabetes"] = 8
    if g("prior_mi") or g("fam_hx"): c["Cardiac history"] = (10 if g("prior_mi") else 0) + (3 if g("fam_hx") else 0)
    if g("prior_pci") or g("prior_cabg"): c["Prior revascularisation"] = 6
    if g("stroke"): c["Stroke / TIA"] = 4
    if g("afib") or "fibrillation" in str(g("rhythm") or "").lower(): c["Atrial fibrillation"] = 4
    if g("lvef") is not None: c["LV function"] = 10 if g("lvef") < 40 else 5 if g("lvef") < 50 else 0
    if g("lvh"): c["LV hypertrophy"] = 3
    if g("wma"): c["Wall-motion abnormality"] = 5
    if g("cac") is not None: c["Calcium score"] = band(g("cac"), [(401, 12), (101, 7), (1, 3)])
    if g("fai") is not None: c["Fat attenuation (FAI)"] = 4 if g("fai") > -70.1 else 0
    if g("stenosis") is not None: c["Coronary stenosis"] = band(g("stenosis"), [(70, 12), (50, 7)])
    if g("troponin") is not None: c["Troponin"] = 8 if g("troponin") > 14 else 0
    if g("bnp") is not None: c["BNP / NT-proBNP"] = band(g("bnp"), [(400, 8), (100, 4)])
    if g("hscrp") is not None: c["hs-CRP"] = 3 if g("hscrp") > 3 else 0
    if g("egfr") is not None: c["Kidney function"] = 4 if g("egfr") < 60 else 0
    if g("ckd"): c["Kidney function"] = 4
    if g("chest_pain"): c["Chest pain"] = 4
    if g("dyspnea"): c["Breathlessness"] = 2
    if g("psych_stress") == "High": c["Psychosocial stress"] = 3
    if g("abd_obesity"): c["Abdominal obesity"] = 2
    if g("activity") == "Low": c["Physical inactivity"] = 3
    if (g("bmi") or 0) >= 30: c["Obesity"] = 3
    c = {k: x for k, x in c.items() if x}
    score = min(100, round(sum(c.values()) * 0.75))  # points are scaled so typical high-risk patients land at 60-85
    tier = "Low Risk" if score <= 40 else "Moderate Risk" if score <= 60 else "High Risk" if score <= 80 else "Critical"
    return score, tier, c, round(100 * sum(v.get(k) is not None for k in CORE) / len(CORE))

# rule-based factors whose inputs are already inside the trained model (so they are not counted twice)
MODEL_COVERED = {"Blood pressure", "LDL cholesterol", "Smoking", "Age", "Diabetes", "Chest pain", "Breathlessness", "LV function", "Obesity"}

def risk(v):
    """(score 0-100, tier, {factor: points}, completeness %).
    The score starts from the CAD probability of the trained model (fields that exist in the Z-Alizadeh Sani dataset).
    Rule-based points for parameters outside the dataset then raise it towards 100: score = m + (100 - m) * e / 100,
    where m is the model probability in % and e the scaled rule points. Falls back to the pure rule score if the model
    is unavailable or fewer than MIN_INPUTS measured inputs are present. Research prototype - NOT a validated clinical tool."""
    base = rule_risk(v)
    pred = cad_model_predict(derive(v))
    if not pred or pred["prob"] is None:
        return base
    extra = {k: p for k, p in base[2].items() if k not in MODEL_COVERED and k != "Cardiac history"}
    if v.get("prior_mi"): extra["Prior heart attack (MI)"] = 10
    m, e = pred["prob"] * 100, min(100, sum(extra.values()) * 0.75)
    score = min(100, round(m + (100 - m) * e / 100))
    tier = "Low Risk" if score <= 40 else "Moderate Risk" if score <= 60 else "High Risk" if score <= 80 else "Critical"
    drivers = {k: round(x * 100) for k, x in pred["drivers"].items() if x * 100 >= 1}
    drivers.update({k: round(p * 0.75) for k, p in extra.items() if round(p * 0.75) >= 1})
    c = dict(sorted(drivers.items(), key=lambda x: -x[1])[:9])
    return score, tier, c, base[3]

def week_forecast(v):
    """CAD risk score for today and the next 6 days. Lab values and history stay as reported; heart rate, blood pressure
    and HRV are varied by a typical day-to-day amount (fixed per patient, so the chart does not jump on every refresh)."""
    import zlib
    seed = zlib.crc32(repr(sorted((k, x) for k, x in v.items() if x is not None and not k.startswith("_"))).encode())
    rng, lim = np.random.RandomState(seed), {p[0]: (p[3], p[4]) for p in PARAMS}
    out = [risk(v)[0]]
    for _ in range(6):
        w = dict(v)
        for k, sd in (("hr", 3.0), ("sbp", 4.0), ("dbp", 2.5)):
            if v.get(k) is not None: w[k] = round(min(lim[k][1], max(lim[k][0], v[k] + rng.normal(0, sd))))
        if v.get("hrv") is not None: w["hrv"] = round(min(lim["hrv"][1], max(lim["hrv"][0], v["hrv"] * (1 + rng.normal(0, 0.08)))))
        out.append(risk(w)[0])
    return out


def _htn(d): return bool(d.get("htn")) or (d.get("sbp") or 0) >= 130 or (d.get("dbp") or 0) >= 80
def _dm(d): return bool(d.get("diabetes")) or (d.get("hba1c") or 0) >= 6.5
def _r(*pairs): return [t for ok, t in pairs if ok]
RULES = [  # ordered: first matching rule decides the scenario
 ("CVD — Heart Failure (HFrEF)", lambda d: _r((d.get("lvef") is not None and d["lvef"] < 40, f"LV ejection fraction {d.get('lvef')}% is below 40%"),
        ((d.get("bnp") or 0) >= 400 and bool(d.get("dyspnea")), f"BNP / NT-proBNP {d.get('bnp')} pg/mL with breathlessness"))),
 ("CVD — Atrial Fibrillation", lambda d: _r((bool(d.get("afib")), "Atrial fibrillation documented"),
        ("fibrillation" in str(d.get("rhythm") or "").lower(), f"ECG rhythm: {d.get('rhythm')}"))),
 ("CVD — Stroke (Cerebrovascular)", lambda d: _r((bool(d.get("stroke")), "Stroke / TIA documented (type cannot be told from values; confirm with imaging)"))),
 ("CAD — Post-MI Recovery", lambda d: _r((bool(d.get("prior_mi")), "Prior heart attack (MI)"), (bool(d.get("prior_pci")), "Prior PCI / stent"), (bool(d.get("prior_cabg")), "Prior bypass surgery (CABG)"))),
 ("CAD — Diabetic + Hypertensive", lambda d: _r((True, "Diabetes (history or HbA1c 6.5% or higher)"), (True, "Hypertension (history or BP 130/80 or higher)")) if _dm(d) and _htn(d) else []),
 ("CAD — Smoker + Sedentary", lambda d: _r((True, "Current smoker"), (True, "Low physical activity")) if d.get("smoking") == "Current" and d.get("activity") == "Low" else []),
 ("CAD — Cardiac Concern", lambda d: _r(((d.get("qtc") or 0) > 470, f"QTc {d.get('qtc')} ms is prolonged"), (abs(d.get("st_mv") or 0) >= 0.1, f"ST deviation {d.get('st_mv')} mV"),
        ((d.get("troponin") or 0) > 14, f"Troponin {d.get('troponin')} ng/L is raised"), (bool(d.get("chest_pain")), "Chest pain / angina reported"),
        ((d.get("stenosis") or 0) >= 50, f"Coronary stenosis {d.get('stenosis')}%"), ((d.get("cac") or 0) > 100, f"Calcium score {d.get('cac')} AU"))),
 ("CAD — High Stress / Low HRV", lambda d: _r((d.get("hrv") is not None and d["hrv"] < 25, f"HRV {d.get('hrv')} ms is low (below 25)"))),
 ("Healthy — Post-Exercise", lambda d: _r(((d.get("hr") or 0) >= 100 and rule_risk(d)[0] <= 40, f"Heart rate {d.get('hr')} bpm with otherwise low risk (recovery pattern)"))),
 ("CAD — Borderline Hypertensive", lambda d: _r((_htn(d), f"Blood pressure {d.get('sbp')}/{d.get('dbp')} mmHg or hypertension history"))),
]

def detect_scenario(v):
    """Pick the best-matching clinical scenario from the patient's own values. Returns None if there is too little data."""
    if sum(v.get(k) is not None for k in CORE) < 3 and not any(v.get(k) for k in FLAGLAB): return None
    d, (score, _, contrib, _) = derive(v), rule_risk(v)
    for name, test in RULES:
        why = test(d)
        if why: break
    else:
        name = "Healthy — Baseline" if score <= 25 else "CAD — Borderline Hypertensive"
        why = ["No values outside normal limits were found"] if score <= 25 else [f"{k}: {p} pts" for k, p in sorted(contrib.items(), key=lambda x: -x[1])[:3]]
    g, n, desc = next((x[0], x[1], x[2]) for x in SCENARIOS if x[1] == name)
    return {"name": name, "group": g, "desc": desc, "why": why}


# ======================= STREAMLIT APP (light mode) =======================
ss = st.session_state
_LOGO_CSS = "[data-testid=stSidebarHeader]{height:auto!important;min-height:7.5rem;padding:1.1rem 1rem .4rem!important;align-items:center}[data-testid=stSidebarHeader] img{height:5.6rem!important;max-height:none!important;width:auto!important;max-width:80%!important}"
st.markdown("""<style>
.stApp{background:linear-gradient(180deg,#eef2f7,#f8fafc 40%)}[data-testid=stSidebar]{background:#fff;border-right:1px solid #e2e8f0}
header[data-testid=stHeader]{background:transparent}.block-container{padding-top:1rem;max-width:1450px}
.card{background:linear-gradient(180deg,#fff,#f8fafc);border:1px solid #e2e8f0;border-top:3px solid #2563eb;border-radius:16px;padding:16px 18px;margin-bottom:16px;
 box-shadow:0 1px 1px rgba(15,23,42,.04),0 4px 8px rgba(15,23,42,.06),0 14px 28px rgba(15,23,42,.08);transition:transform .2s,box-shadow .2s}
.card:hover{transform:translateY(-3px);box-shadow:0 2px 2px rgba(15,23,42,.05),0 8px 16px rgba(15,23,42,.09),0 22px 40px rgba(15,23,42,.13)}
.card h6{color:#475569;font-size:.78rem;letter-spacing:.06em;margin:0 0 10px;text-transform:uppercase;font-weight:700}
.big{font-size:2.3rem;font-weight:800;color:#0f172a;line-height:1.1}.u{color:#64748b;font-size:1rem;font-weight:500}
.pill{display:inline-block;padding:2px 10px;border-radius:999px;font-size:.74rem;font-weight:700;border:1.5px solid;background:#fff}
.row{display:flex;justify-content:space-between;align-items:center;border-bottom:1px solid #e2e8f0;padding:8px 0;font-size:.95rem}
.row:last-child{border-bottom:0}.mut{color:#64748b;font-size:.82rem}
.chip{display:inline-block;background:#fee2e2;color:#991b1b;border-radius:999px;padding:3px 11px;margin:3px 4px 3px 0;font-size:.82rem;font-weight:600}
[data-testid=stVerticalBlockBorderWrapper]:has(.brand){background:#fff;border:1px solid #e2e8f0;border-radius:14px;box-shadow:0 4px 14px rgba(15,23,42,.07);padding:4px 8px}
.brand{font-size:1.15rem;font-weight:800;color:#0f172a;line-height:1.25}.brand .page{display:block;color:#2563eb;font-weight:600;font-size:.9rem}
.who{font-size:.9rem;line-height:1.35}
.shield{width:84px;height:84px;border-radius:50%;background:#dbeafe;border:1px solid #bfdbfe;display:grid;place-items:center;margin:auto;box-shadow:0 10px 24px rgba(37,99,235,.22)}
.hero{margin:14px 0 6px;font-size:2.1rem;font-weight:800;color:#0f172a}.sub{color:#64748b;max-width:640px;margin:auto;font-size:1.05rem}
[data-testid=stExpander],[data-testid=stForm]{background:#fff;border:1px solid #e2e8f0!important;border-radius:16px;box-shadow:0 6px 18px rgba(15,23,42,.07)}
[data-testid=stToolbar],[data-testid=stDecoration]{display:none}
[data-testid=stFileUploaderDropzone]{border:2px dashed #94a3b8;border-radius:16px;background:#fff;padding:28px 22px;box-shadow:0 8px 24px rgba(15,23,42,.07)}
[data-testid=stFileUploaderDropzone]:hover{border-color:#2563eb;background:#f8fbff}
[data-testid=stFileUploaderDropzone] button{background:#0f172a!important;border:0!important;border-radius:10px;min-height:42px;padding:0 22px}
[data-testid=stFileUploaderDropzone] button *{color:#fff!important}
[data-testid^=stBaseButton-primary]{background:#0f172a!important;border:1px solid #0f172a!important;border-radius:10px;min-height:42px}
[data-testid^=stBaseButton-primary] *{color:#fff!important}
[data-testid^=stBaseButton-secondary],[data-testid=stPopoverButton]{background:#fff!important;border:1px solid #94a3b8!important;border-radius:10px;min-height:42px}
[data-testid^=stBaseButton-secondary] *,[data-testid=stPopoverButton] *{color:#0f172a!important;font-weight:600}
[data-testid^=stBaseButton]:hover,[data-testid=stPopoverButton]:hover{border-color:#2563eb!important}
[data-testid=stVerticalBlockBorderWrapper]:has(.brand) [data-testid=stHorizontalBlock]{align-items:center}
[data-testid=stSidebarHeader]{height:2.4rem!important;min-height:0!important;padding:.4rem .8rem!important}
[data-testid=stSidebar] [data-testid=stImage]{margin:-.4rem 0 .6rem}
[data-testid=stVerticalBlockBorderWrapper]:not(:has(.brand)){background:#fff;border:1px solid #e2e8f0;border-radius:14px;box-shadow:0 4px 12px rgba(15,23,42,.05)}
.plabel{font-weight:700;font-size:.8rem;letter-spacing:.05em;color:#334155;text-transform:uppercase}
.miss{color:#b45309;font-style:italic;font-size:.86rem}
[data-testid=stExpander] summary{background:#f1f5f9;border-radius:12px;font-weight:800;letter-spacing:.07em;text-transform:uppercase;font-size:.85rem}
.stTabs [data-baseweb=tab]{font-size:1rem;font-weight:700;padding:10px 18px}
[data-testid=stNumberInput] input,[data-testid=stTextInput] input{border-radius:10px}
.card{overflow-wrap:anywhere;min-width:0}[data-testid=stColumn],[data-testid=column]{min-width:0}
div[data-testid=stElementContainer]:has(.sidemark){display:none}
[data-testid=stColumn]:has(.sidemark),[data-testid=column]:has(.sidemark){position:sticky;top:1rem;align-self:flex-start;max-height:calc(100vh - 2rem);overflow-y:auto;overflow-x:hidden;scrollbar-width:thin;padding:4px 10px 10px 4px}
</style>""", unsafe_allow_html=True)
G, A, R, B = "#16a34a", "#d97706", "#dc2626", "#2563eb"
H = lambda s: st.markdown("".join(l.strip() for l in s.splitlines()), unsafe_allow_html=True)
fmt = lambda x, u="": "—" if x is None else f"{x:g}{u}" if isinstance(x, (int, float)) and not isinstance(x, bool) else str(x)
pill = lambda t, c: f'<span class="pill" style="color:{c};border-color:{c}">{t}</span>'
tcol = lambda s: G if s <= 40 else A if s <= 60 else R


def clear_all():
    ss.files, ss.seen, ss.up, ss.conflicts = [], set(), ss.get("up", 0) + 1, []
    ss.mode, ss.scn, ss.base, ss.manual, ss.src = "empty", None, dict(MISSING), {}, {}

def load_scenario(name):
    ss.mode, ss.scn, ss.base, ss.manual, ss.src = "scenario", name, scenario_values(name), {}, {}

def rebuild():
    merged, ss.src, ss.conflicts = {}, {}, []
    for f in ss.files:
        for k, v in f["params"].items():
            if k in merged and merged[k] != v: ss.conflicts.append((k, merged[k], ss.src[k], v, f["file"]))
            merged[k], ss.src[k] = v, f["file"]
    ss.mode, ss.base = "report", {**MISSING, **merged}
    ss.manual = {k: v for k, v in ss.manual.items() if k not in merged}

if "mode" not in ss: clear_all()
get = lambda k: ss.manual[k] if k in ss.manual else ss.base.get(k)
values = lambda: {k: get(k) for k in KEYS}

def uploader(label="Upload one or more patient reports (PDF, photo/scan, DOCX, XLSX, CSV, JSON, TXT)"):
    files = st.file_uploader(label, type=EXT, accept_multiple_files=True, key=f"up{ss.up}", label_visibility="collapsed")
    new = [f for f in files if (f.name, f.size) not in ss.seen]
    if new:
        with st.spinner(f"Reading {len(new)} file(s)…"):
            for f in new:
                ss.seen.add((f.name, f.size)); ss.files.append(extract(f.name, f.getvalue()))
        rebuild(); st.rerun()

def gauge(score, tier):
    col, L = tcol(score), 0.75 * 2 * 3.14159 * 60
    return (f'<svg viewBox="0 0 160 160" width="100%"><g transform="rotate(135 80 80)">'
            f'<circle cx="80" cy="80" r="60" fill="none" stroke="#e2e8f0" stroke-width="11" stroke-dasharray="{L} 999" stroke-linecap="round"/>'
            f'<circle cx="80" cy="80" r="60" fill="none" stroke="{col}" stroke-width="11" stroke-linecap="round" stroke-dasharray="{L * score / 100} 999"/></g>'
            f'<text x="80" y="88" fill="#0f172a" font-size="40" font-weight="800" text-anchor="middle">{score}</text>'
            f'<text x="80" y="107" fill="{col}" font-size="10" text-anchor="middle" font-weight="700">{tier.upper()}</text></svg>')

def donut(c):
    pal = ["#dc2626", "#d97706", "#16a34a", "#7c3aed", "#0d9488", "#ea580c", "#4f46e5", "#db2777", "#ca8a04"]
    tot, off, out = sum(c.values()) or 1, 0, ""
    for i, (k, p) in enumerate(c.items()):
        out += f'<circle cx="60" cy="60" r="40" fill="none" stroke="{pal[i % 9]}" stroke-width="16" stroke-dasharray="{251.3 * p / tot} 999" stroke-dashoffset="{-off}" transform="rotate(-90 60 60)"/>'
        off += 251.3 * p / tot
    return f'<svg viewBox="0 0 120 120" width="55%" style="display:block;margin:auto">{out}<text x="60" y="65" fill="#0f172a" text-anchor="middle" font-size="16" font-weight="800">{sum(c.values())}</text></svg>', pal

def model_card(v):
    """Which trained model produced the score, and how every candidate scored in cross-validation."""
    M = train_cad_model()
    if not M["ok"]:
        H(f'<div class="card"><h6>Prediction model</h6><div style="font-size:.88rem">The trained model is not available ({M["error"]}). The rule-based score is shown instead.</div></div>')
        return
    pred = cad_model_predict(derive(v))
    used = (f'{pred["present"]} of {len(MEASURED)} measured inputs available; missing ones use typical training values, and findings not reported count as absent.'
            if pred and pred["prob"] is not None else
            f'Only {pred["present"]} of {len(MEASURED)} measured inputs available (at least {MIN_INPUTS} are needed), so the rule-based score is shown.')
    best = next(r for r in M["cv"] if r[0] == M["best"])
    rows = "".join(f'<div class="row" style="padding:3px 0;font-size:.82rem;{"font-weight:800" if n == M["best"] else "color:#64748b"}"><span>{"✓ " if n == M["best"] else ""}{n}</span><span>{m:.3f}</span></div>' for n, m, s in M["cv"])
    H(f'<div class="card"><h6>Prediction model</h6><div style="font-size:.95rem;font-weight:800">{M["best"]}</div>'
      f'<div class="mut" style="margin:2px 0 8px">Highest cross-validated ROC-AUC: {best[1]:.3f} ± {best[2]:.3f} · held-out test ROC-AUC: {M["test_auc"]:.3f}</div>'
      f'<div style="font-size:.84rem;margin-bottom:8px">Trained on {M["n"]} patients (Z-Alizadeh Sani dataset) to predict CAD confirmed on angiography. Parameters outside the dataset (troponin, calcium score, QTc, HRV and others) are rule-based and added on top.</div>'
      f'<div class="mut" style="margin-top:4px">{used}</div></div>')
    with st.expander("See all models compared"):
        H(f'<div class="plabel" style="margin:0 0 2px">Cross-validated ROC-AUC</div>{rows}')

def trend_card(v):
    sc = week_forecast(v); tier = risk(v)[1]
    delta = sc[-1] - sc[0]
    word, arrow = ("Stable", "→") if abs(delta) <= 2 else ("Rising", "↑") if delta > 0 else ("Falling", "↓")
    lo, hi = min(sc), max(sc); span = max(hi - lo, 6)
    xs = [40 + i * 620 / 6 for i in range(7)]
    ys = [75 - (s - (lo + hi) / 2) / span * 50 for s in sc]
    line = " ".join(f"{x:.1f},{y:.1f}" for x, y in zip(xs, ys))
    pts = "".join(f'<circle cx="{x:.1f}" cy="{y:.1f}" r="5" fill="{tcol(s)}"/><text x="{x:.1f}" y="{y - 13:.1f}" fill="#0f172a" font-size="14" font-weight="800" text-anchor="middle">{s}</text>'
                  f'<text x="{x:.1f}" y="138" fill="#64748b" font-size="12" text-anchor="middle">D{i + 1}</text>' for i, (x, y, s) in enumerate(zip(xs, ys, sc)))
    H(f'<div class="card"><h6>7-day CAD risk trend (projected)</h6>'
      f'<div style="display:flex;justify-content:space-between;align-items:baseline;flex-wrap:wrap;gap:6px"><div><span class="big">{sc[0]}</span> {pill(tier, tcol(sc[0]))}</div>'
      f'<div class="mut" style="font-size:.95rem;font-weight:700">{sc[0]} → {sc[-1]} &nbsp;{arrow} {word}</div></div>'
      f'<svg viewBox="0 0 700 150" width="100%"><polyline points="{line}" fill="none" stroke="{tcol(sc[0])}" stroke-width="3" stroke-linejoin="round" stroke-linecap="round"/>{pts}</svg>'
      f'<div class="mut">D1 is today\'s score. D2 to D7 re-score the same patient with typical day-to-day variation in heart rate, blood pressure and HRV; lab values and history are held as reported. '
      f'This is a projection, not a forecast from real daily measurements.</div></div>')

LAYMAN = {  # why each factor matters, in everyday words
    "Age": "Heart disease becomes more common as people get older.",
    "BMI": "Extra body weight makes the heart work harder and tends to raise cholesterol and blood sugar.",
    "Obesity": "Extra body weight makes the heart work harder and tends to raise cholesterol and blood sugar.",
    "Diabetes": "High blood sugar over the years damages the walls of the heart's arteries.",
    "Hypertension": "High blood pressure strains the heart and slowly damages the arteries.",
    "Blood pressure": "High blood pressure strains the heart and slowly damages the arteries.",
    "Systolic BP": "A high top blood-pressure number strains the heart and the arteries.",
    "Current smoking": "Smoking damages artery walls and lets blockages build up faster.",
    "Smoking": "Smoking damages artery walls and lets blockages build up faster.",
    "Family history of CAD": "A close relative with heart disease means it may run in the family.",
    "Chest pain / angina": "Chest pain can be a sign that the heart is not getting enough blood.",
    "Chest pain": "Chest pain can be a sign that the heart is not getting enough blood.",
    "Breathlessness": "Getting out of breath easily can be a sign that the heart is under strain.",
    "Heart rate": "A fast resting pulse means the heart is working harder than usual.",
    "LDL cholesterol": "LDL is the 'bad' cholesterol that can build up inside arteries and block them.",
    "HDL cholesterol": "HDL is the 'good' cholesterol. A low level means less protection for the arteries.",
    "Triglycerides": "High levels of this blood fat are linked to clogged arteries.",
    "Total cholesterol": "High overall cholesterol raises the chance of blocked arteries.",
    "ApoB": "ApoB counts the cholesterol particles that can clog arteries.",
    "Lipoprotein(a)": "Lp(a) is an inherited type of cholesterol that raises heart risk.",
    "Hemoglobin": "Hemoglobin carries oxygen. A level outside the normal range makes the heart work harder.",
    "LV ejection fraction": "This shows how well the heart pumps. A lower value means a weaker pump.",
    "LV function": "This shows how well the heart pumps. A lower value means a weaker pump.",
    "Low HRV": "Low heart-rate variability is a sign that the body is under stress.",
    "Physiological stress": "Signs of stress on the body, such as a fast pulse, add strain on the heart.",
    "QTc interval": "A longer-than-normal QT time on the ECG means the heart's electrical reset is slow.",
    "QRS duration": "A wider QRS on the ECG means the heart's electrical signal is travelling slowly.",
    "PR interval": "A long PR interval on the ECG means the heart's electrical signal is delayed.",
    "ST deviation": "A shift in the ST part of the ECG can point to reduced blood flow to the heart.",
    "Prior heart attack (MI)": "Having had a heart attack raises the chance of another one.",
    "Cardiac history": "Past heart problems or a family history raise the chance of more.",
    "Prior revascularisation": "A past stent or bypass shows the heart's arteries have been blocked before.",
    "Stroke / TIA": "A past stroke shows the blood vessels are already affected.",
    "Atrial fibrillation": "An irregular heartbeat can strain the heart and raise stroke risk.",
    "LV hypertrophy": "A thickened heart wall usually means the heart has been working too hard for a long time.",
    "Wall-motion abnormality": "Part of the heart muscle is not moving as well as it should.",
    "Calcium score": "A calcium score above zero means hardened plaque has been found in the heart's arteries.",
    "Fat attenuation (FAI)": "This scan marker points to inflammation around the heart's arteries.",
    "Coronary stenosis": "A narrowing of the heart's arteries has been seen on imaging.",
    "Troponin": "A raised troponin can mean the heart muscle is being injured.",
    "BNP / NT-proBNP": "A raised BNP is a sign that the heart is under strain.",
    "hs-CRP": "A raised hs-CRP points to inflammation in the body, which is linked to artery disease.",
    "Kidney function": "Weak kidneys and heart disease often go together.",
    "Psychosocial stress": "Long-term stress can raise blood pressure and strain the heart.",
    "Abdominal obesity": "Fat around the belly is linked to higher heart risk.",
    "Physical inactivity": "Not moving enough weakens the heart and raises blood pressure and weight.",
}
# tier -> (headline, what it means, what the patient can do, what caregivers can do, suggested next step)
TIER_TEXT = {
    "Low Risk": ("Low risk: the numbers look reassuring",
                 "On the information given, a significant blockage in the heart's arteries looks unlikely. This is an estimate, not a guarantee.",
                 "Keep up healthy habits: regular activity, a balanced diet, no smoking, and routine check-ups.",
                 "Encourage the healthy routine and make sure the patient reports any new symptoms.",
                 "A routine check-up with the doctor is enough. Repeat screening as the doctor advises."),
    "Moderate Risk": ("Moderate risk: worth discussing with a doctor",
                      "Some factors are raising the risk. This does not mean heart disease is present, but it is worth looking into.",
                      "Book a visit with your doctor in the coming weeks and bring these results. Work on the factors listed below.",
                      "Help with medicines, healthy meals and activity. Watch for chest discomfort or breathlessness.",
                      "See a doctor in the next few weeks. They may suggest simple tests such as an ECG or a stress test."),
    "High Risk": ("High risk: please see a doctor soon",
                  "Several factors together point to a higher chance of blocked heart arteries, even if you feel well.",
                  "Do not wait for symptoms. Arrange a doctor or heart-specialist visit soon and check with them before heavy exertion.",
                  "Help arrange the appointment, go along, and bring all reports and medicine lists. Learn the warning signs below.",
                  "Arrange a doctor or cardiologist visit within days to a couple of weeks. Tests to look at the arteries may be advised."),
    "Critical": ("Very high risk: get a medical review promptly",
                 "The numbers point strongly to heart-artery disease, so this needs a doctor's attention.",
                 "Contact your doctor or a heart specialist as soon as possible, ideally within a day or two. Avoid heavy exertion until you have been seen.",
                 "Help the patient get seen quickly and stay with them if possible. Learn the warning signs below and act on them at once.",
                 "Get a medical review as soon as possible. If any warning sign below appears, do not wait."),
}
LADDER = [("Low", "0–40", "Healthy habits and routine check-ups are enough."), ("Moderate", "41–60", "Talk to a doctor in the coming weeks."),
          ("High", "61–80", "See a doctor soon, even without symptoms."), ("Very high", "81–100", "Get a medical review as soon as possible.")]

def layman_card(score, tier, c):
    """Plain-language reading of the score for patients and caregivers (stacked, sized for the right-hand column)."""
    head, meaning, pat, care, nxt = TIER_TEXT.get(tier, TIER_TEXT["Moderate Risk"]); col = tcol(score)
    why = "".join(f'<div style="margin:0 0 8px;font-size:.86rem"><b>{k}</b><br><span style="color:#475569">{LAYMAN[k]}</span></div>'
                  for k in [k for k in sorted(c, key=c.get, reverse=True) if k in LAYMAN][:3])
    lab = lambda t: f'<div class="plabel" style="margin:14px 0 3px">{t}</div>'
    H(f'<div class="card"><h6>Your score in plain words</h6>'
      f'<div style="font-size:1.05rem;font-weight:800;color:{col};margin-bottom:4px;line-height:1.3">{head}</div>'
      f'<div style="font-size:.88rem">Score <b>{score} out of 100</b>. {meaning}</div>'
      f'{lab("What is raising the score" if score > 40 else "Small factors in the score")}{why or "<div class=mut>No single factor stands out.</div>"}'
      f'{lab("For the patient")}<div style="font-size:.86rem">{pat}</div>'
      f'{lab("For caregivers and family")}<div style="font-size:.86rem">{care}</div>'
      f'{lab("Suggested next step")}<div style="font-size:.86rem"><b>{nxt}</b></div>'
      f'<div style="background:#fef2f2;border:1px solid #fecaca;color:#991b1b;border-radius:10px;padding:10px 12px;margin-top:14px;font-size:.82rem;line-height:1.45">'
      f'<b>Call your local emergency number right away</b> for chest pain, pressure or tightness lasting more than a few minutes, pain spreading to the arm, jaw, neck or back, '
      f'sudden breathlessness, fainting, or a cold sweat with nausea. This applies at every risk level.</div>'
      f'<div class="mut" style="margin-top:8px">A screening estimate from a computer model, not a diagnosis. Only a doctor can confirm heart disease.</div></div>')

def side(v):
    st.markdown('<div class="sidemark"></div>', unsafe_allow_html=True)
    score, tier, c, conf = risk(v)
    bands = [(0, 20, "Very low"), (21, 40, "Low"), (41, 60, "Moderate"), (61, 80, "High"), (81, 100, "Critical")]
    rows = "".join(f'<div class="row" style="padding:3px 0;{"font-weight:800" if a <= score <= b else "color:#64748b"}"><span>{a}–{b}</span><span>{n}</span></div>' for a, b, n in bands)
    H(f'<div class="card"><h6>Coronary artery disease risk score</h6>{gauge(score, tier)}<div class="mut" style="text-align:center">Data completeness: {conf}%</div></div>')
    layman_card(score, tier, c)
    model_card(v)
    H(f'<div class="card"><h6>What this means</h6><div style="font-size:.88rem;margin-bottom:8px">Based on {"your uploaded report values" if ss.mode == "report" else "the values you entered" if ss.mode == "manual" else "simulated demo values"}: <b>{tier}</b>. Trained-model screening score with rule-based additions — not a diagnosis.</div>{rows}</div>')
    if c:
        svg, pal = donut(c)
        leg = "".join(f'<div class="row"><span><span style="color:{pal[i % 9]}">●</span> {k}</span><b>{p} pts</b></div>' for i, (k, p) in enumerate(sorted(c.items(), key=lambda x: -x[1])))
        H(f'<div class="card"><h6>What drives the score</h6>{svg}{leg}<div class="mut">Points show how much each factor raised the estimate; they are approximate and do not add up exactly to the score.</div></div>')

def page(title, allow_empty=False):
    topbar(title)
    if ss.mode == "scenario":
        st.warning("Demo scenario loaded — simulated values, not a real patient.")
        if st.button("Exit demo and clear"): clear_all(); st.rerun()
    if ss.mode == "empty":
        if not allow_empty:
            st.info("No data yet. Upload a patient report to see results on this page.")
            uploader(); st.stop()
        return st.container(), derive(values())
    v = values(); score, tier, _, _ = risk(v)
    main, right = st.columns([3, 1.1], gap="medium")
    with right: side(derive(v))
    return main, derive(v)

def metric(title, val, unit, status="", col=G, note=""):
    H(f'<div class="card"><h6>{title}</h6><div class="big">{val} <span class="u">{unit}</span></div><div style="color:{col};font-weight:700;margin-top:4px">{status}</div><div class="mut">{note}</div></div>')

def stat(k, v):
    lo, hi = REF.get(k, (None, None))
    if lo is not None and v < lo: return "Low", A
    if hi is not None and v > hi: return "High", R if k in ("troponin", "stenosis", "bnp", "cac") else A
    return "Normal", G

def miss_label():
    return "Not present in report" if ss.mode == "report" else "Not entered" if ss.mode == "manual" else "Not in demo data"

def card(title, keys, d, extra=""):
    rows, have = extra, 0
    for k in keys:
        v = d.get(k)
        if v is None:
            rows += f'<div class="row"><span>{LAB[k][0]}</span><span class="miss">{miss_label()} · enter manually</span></div>'; continue
        have += 1; t, c = stat(k, v); est = ' <span class="mut">(calculated)</span>' if k in d.get("_est", ()) else ""
        rows += f'<div class="row"><span>{LAB[k][0]}{est}</span><span><b>{fmt(v)}</b> <span class="mut">{LAB[k][1]}</span> {pill(t, c)}</span></div>'
    H(f'<div class="card"><h6>{title} <span class="mut" style="text-transform:none">({have} of {len(keys)} present)</span></h6>{rows}</div>')

def field(k, label, unit, lo, hi, step):
    ss["w_" + k] = None if get(k) is None else float(get(k))
    st.number_input(f"{label} ({unit})", float(lo), float(hi), value=None, step=float(step), format="%g", key="w_" + k, on_change=lambda: ss.manual.__setitem__(k, ss["w_" + k]))
    if get(k) is None and ss.mode != "empty": st.markdown(f":orange[{miss_label()} · enter manually]")
    elif ss.mode == "report" and k in ss.src: st.caption("edited manually" if k in ss.manual else f"from {ss.src[k]}")

def name_input():
    ss.w_name = get("name") or ""
    st.text_input("Name", key="w_name", on_change=lambda: ss.manual.__setitem__("name", ss.w_name))

def choice(k, label, opts):
    ss["w_" + k] = get(k) if get(k) in opts else None
    st.radio(label, opts, horizontal=True, index=None, key="w_" + k, on_change=lambda: ss.manual.__setitem__(k, ss["w_" + k]))

def check(k):
    ss["w_" + k] = bool(get(k))
    st.checkbox(FLAGLAB[k], key="w_" + k, on_change=lambda: ss.manual.__setitem__(k, ss["w_" + k]))

# ---------- pages ----------
def topbar(title):
    import html as _h
    ov, has = title == "Patient Overview", ss.mode != "empty"
    v = values()
    score, tier = risk(v)[:2] if has else (0, "")
    age = f"{int(v['age'])} yrs" if v.get("age") is not None else None
    who = " · ".join(_h.escape(str(x)) for x in (v.get("name"), age) if x) or ("Patient details not set" if has else "No patient loaded")
    if has and age and not v.get("name"): who = "Name not entered · " + age
    src = (f"{len(ss.files)} uploaded report(s)" if ss.mode == "report" else "Manual entry" if ss.mode == "manual" else f"Demo: {ss.scn}") if has else "Upload a report to begin"
    det = detect_scenario(v) if has else None
    badge = pill("● " + tier, tcol(score)) if has else ""
    with st.container(border=True):
        c = st.columns([3, 3, 3, 2, 1.2], vertical_alignment="center")
        with c[0]: H(f'<div class="brand">Cardio Care AI<span class="page">{title}</span></div>')
        with c[1]: H(f'<div class="who"><b style="font-size:1.12rem">{who}</b><br><span class="mut">{" · ".join(x for x in (v.get("sex"), src) if x)}</span></div>')
        with c[2]: H(f'<div class="who"><span class="mut">Detected scenario</span><br><b>{det["name"] if det else "Awaiting patient data"}</b> {badge}</div>')
        if ov:
            if c[3].button("Open Dashboard", key="tb_dash", type="primary", use_container_width=True, disabled=not has): st.switch_page(DASH)
        elif c[3].button("Patient Overview", key="tb_ov", use_container_width=True): st.switch_page(OVERVIEW)
        if not ov and c[4].button("Upload", key="tb_up", use_container_width=True): st.switch_page(OVERVIEW)

def hint(k):
    lo, hi = REF.get(k, (None, None))
    return f"Normal {lo:g}–{hi:g}" if lo is not None and hi is not None else f"Normal < {hi:g}" if hi is not None else f"Normal > {lo:g}" if lo is not None else ""

def _store(k, wk):
    if ss.mode in ("empty", "scenario"): ss.base, ss.src, ss.manual, ss.mode = dict(MISSING), {}, {}, "manual"
    ss.manual[k] = ss[wk]

def entry(col, k, label, unit, lo, hi, step, ph=""):
    wk = "e_" + k; ss[wk] = None if get(k) is None else float(get(k))
    with col:
        st.number_input(f"{label} ({unit})", float(lo), float(hi), value=None, step=float(step), format="%g", key=wk, placeholder=ph or hint(k), on_change=_store, args=(k, wk))
        v = get(k)
        if k in REF:
            H(f'<div style="margin:-10px 0 10px">{pill(*stat(k, v)) if v is not None else ""} <span class="mut">{hint(k)}</span></div>')
        if v is None and ss.mode == "report": st.markdown(":orange[Not present in report · enter here]")

def seg(col, k, label, opts):
    wk = "e_" + k; ss[wk] = get(k) if get(k) in opts else None
    with col:
        st.caption(label)
        if hasattr(st, "segmented_control"): st.segmented_control(label, opts, key=wk, label_visibility="collapsed", on_change=_store, args=(k, wk))
        else: st.radio(label, opts, horizontal=True, index=None, key=wk, label_visibility="collapsed", on_change=_store, args=(k, wk))

def flag(col, k):
    wk = "e_" + k; ss[wk] = bool(get(k))
    col.checkbox(FLAGLAB[k], key=wk, on_change=_store, args=(k, wk))

def manual_entry():
    if ss.mode == "scenario":
        st.info("Demo data is loaded. Exit the demo to enter patient values manually.")
        if st.button("Exit demo and start manual entry", key="exit_demo_manual"): clear_all(); st.rerun()
        return
    done = risk(values())[3]
    st.markdown("##### Clinical data entry")
    st.caption("Enter what is available; anything left blank is treated as unknown. Each value is checked against its reference range as you type, and the dashboard updates automatically.")
    st.progress(done / 100, text=f"Key clinical values entered: {done}%")
    tabs = st.tabs(["Patient"] + GROUPS + ["History & symptoms"])
    with tabs[0]:
        with st.container(border=True):
            st.markdown("**Demographics**")
            c1, c2, c3 = st.columns(3)
            ss.e_name = get("name") or ""
            c1.text_input("Full name", key="e_name", placeholder="e.g. Ravi Kumar", on_change=_store, args=("name", "e_name"))
            entry(c2, "age", "Age", "years", 1, 120, 1, "e.g. 52"); seg(c3, "sex", "Sex", ["Male", "Female"])
            c1, c2, c3 = st.columns(3)
            entry(c1, "height_cm", "Height", "cm", 100, 230, 1, "e.g. 170"); entry(c2, "weight_kg", "Weight", "kg", 25, 300, 0.5, "e.g. 70")
            bmi = derive(values()).get("bmi")
            c3.metric("BMI (calculated)", f"{bmi:g}" if bmi else "—")
            if bmi: c3.caption("Underweight" if bmi < 18.5 else "Normal weight" if bmi < 25 else "Overweight" if bmi < 30 else "Obese")
        with st.container(border=True):
            st.markdown("**Lifestyle**")
            c1, c2 = st.columns(2)
            seg(c1, "smoking", "Smoking", ["Never", "Former", "Current"]); seg(c2, "activity", "Physical activity", ["Low", "Moderate", "High"])
            seg(c1, "diet", "Diet quality", ["Low", "Moderate", "High"]); seg(c2, "alcohol", "Alcohol use", ["None", "Light", "Moderate", "Heavy"]); seg(c1, "psych_stress", "Psychosocial stress", ["Low", "Moderate", "High"])
    for t, grp in zip(tabs[1:-1], GROUPS):
        with t, st.container(border=True):
            st.markdown(f"**{grp}**")
            cols = st.columns(3)
            for i, p in enumerate(p for p in PARAMS if p[7] == grp): entry(cols[i % 3], p[0], p[1], p[2], p[3], p[4], p[5])
    with tabs[-1]:
        for title, ks in (("Medical history", ["diabetes", "htn", "dyslip", "fam_hx", "abd_obesity", "prior_mi", "prior_pci", "prior_cabg", "stroke", "ckd", "afib"]),
                          ("ECG and imaging findings", ["lvh", "wma", "lbbb"]), ("Symptoms", ["chest_pain", "dyspnea", "fatigue", "palp", "syncope"])):
            with st.container(border=True):
                st.markdown(f"**{title}**")
                cols = st.columns(3)
                for i, k in enumerate(ks): flag(cols[i % 3], k)
    a, b, _ = st.columns([1.6, 1.4, 3])
    if a.button("View dashboard", key="man_dash", type="primary", use_container_width=True, disabled=ss.mode == "empty"): st.switch_page(DASH)
    if b.button("Clear all data", key="man_clear", use_container_width=True): clear_all(); st.rerun()

def p_overview():
    main, d = page("Patient Overview", allow_empty=True)
    with main:
        H('<div style="text-align:center;margin:20px 0 14px"><div class="shield"><svg width="40" height="40" viewBox="0 0 24 24" fill="none" stroke="#2563eb" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="M12 3l7 3v5c0 5-3 8-7 10-4-2-7-5-7-10V6z"/><path d="M9 12l2 2 4-4"/></svg></div>'
          '<div style="font-size:2rem;font-weight:800;color:#0f172a;margin:16px 0 6px;line-height:1.2">Patient Overview</div>'
          '<div style="max-width:640px;margin:0 auto;font-size:1.02rem;color:#64748b;line-height:1.5">Upload a clinical or laboratory report to view your cardiovascular health assessment, risk scores, trends, and personalized recommendations.</div></div>')
        t1, t2 = st.tabs(["Upload reports", "Enter manually"])
        with t1:
            _, mid, _ = st.columns([1, 3.2, 1])
            with mid:
                H('<div style="text-align:center;font-weight:700;font-size:1.1rem;margin-bottom:2px">Click Upload or drag and drop to add clinical reports</div><div class="mut" style="text-align:center;margin-bottom:10px">PDF, JPG/PNG scans, DOCX, XLSX, CSV, JSON or TXT. You can select several files at once.</div>')
                uploader()
                H('<div class="mut" style="text-align:center;margin-top:8px">Private · files are read in memory and are not stored</div>')
            if ss.files:
                st.success(f"{len(ss.files)} report(s) processed · {sum(len(f['params']) for f in ss.files)} values found")
                for f in ss.files:
                    st.markdown(f"**{f['file']}** — {len(f['params'])} values" + "".join(f"  \n:orange[{n}]" for n in f["notes"]))
                    if "name" not in f["params"] and f.get("text"):
                        with st.expander(f"Patient name not found in {f['file']}: show the text read from it"): st.code(f["text"][:1500], language=None)
                missing = [LAB[k][0] for k in LAB if get(k) is None]
                if missing:
                    with st.expander(f"{len(missing)} parameters are not present in the uploaded report(s)"):
                        st.write(", ".join(missing)); st.caption("Enter them in the Enter manually tab, on Lab Report, or on Patient Profile.")
                miss = [l for k, l in (("name", "name"), ("age", "age"), ("sex", "sex")) if get(k) in (None, "")]
                if miss:
                    st.warning("Not found in the reports: " + ", ".join(miss) + ". Enter them here:")
                    a, b, c = st.columns(3)
                    with a: name_input()
                    with b: field("age", "Age", "yrs", 1, 120, 1)
                    with c: choice("sex", "Sex", ["Male", "Female"])
                if st.button("Clear all reports"): clear_all(); st.rerun()
        with t2: manual_entry()
        if ss.mode == "empty":
            H('<div class="mut" style="text-align:center;margin-top:22px">No report handy? Explore the app with sample data</div>')
            _, b, c, _ = st.columns([2, 1.6, 1.6, 2])
            if b.button("Healthy demo", use_container_width=True): load_scenario("Healthy — Baseline"); st.switch_page(DASH)
            if c.button("Heart-disease (CAD) demo", use_container_width=True): load_scenario("CAD — Borderline Hypertensive"); st.switch_page(DASH)

def scenario_card(v):
    det = detect_scenario(v)
    if not det: return
    score = risk(v)[0]
    why = "".join(f'<div class="row"><span>{w}</span></div>' for w in det["why"])
    grp = {"Healthy": "Healthy", "CAD": "Coronary artery disease", "Other CVD": "Other cardiovascular disease"}[det["group"]]
    H(f'<div class="card"><h6>Scenario analysis (detected from the patient data)</h6><div style="font-size:1.3rem;font-weight:800;margin-bottom:2px">{det["name"]} {pill(grp, tcol(score))}</div>'
      f'<div class="mut" style="margin-bottom:8px">{det["desc"]}</div><div style="font-weight:700;font-size:.85rem;margin-bottom:2px">Why this scenario</div>{why}'
      f'<div class="mut" style="margin-top:8px">Rule-based screening aid. It does not replace a clinician\'s diagnosis.</div></div>')

def p_dashboard():
    main, d = page("Dashboard")
    with main:
        scenario_card(values())
        trend_card(values())
        miss = [LAB[k][0] for k in LAB if d.get(k) is None]
        if miss and ss.mode != "scenario":
            st.warning(f"{len(miss)} of {len(LAB)} parameters are {'not present in the uploaded report(s)' if ss.mode == 'report' else 'not entered yet'}. They are marked below. Enter them manually on the Lab Report or Patient Profile page.")
            a, b, _ = st.columns([1.3, 1.5, 3]); a.page_link(LABP, label="Open Lab Report"); b.page_link(PROFILE, label="Open Patient Profile")
        s, dia = d.get("sbp"), d.get("dbp"); t, tc = bp_class(s, dia); ml = miss_label()
        absent = lambda title, unit: (title, "—", unit, ml, A, "Enter it manually")
        tiles = [("Heart rate", fmt(d["hr"]), "bpm", *stat("hr", d["hr"]), "Normal 60–100 bpm") if d.get("hr") is not None else absent("Heart rate", "bpm"),
                 ("Blood pressure", f"{fmt(s)}/{fmt(dia)}", "mmHg", t, tc, "Normal < 120/80") if s is not None else absent("Blood pressure", "mmHg"),
                 ("LV ejection fraction", fmt(d["lvef"]), "%", *stat("lvef", d["lvef"]), "Normal 50–75%") if d.get("lvef") is not None else absent("LV ejection fraction", "%"),
                 ("HRV (RMSSD)", fmt(d["hrv"]), "ms", *stat("hrv", d["hrv"]), "Normal above 30 ms") if d.get("hrv") is not None else absent("HRV (RMSSD)", "ms")]
        for col, tl in zip(st.columns(4), tiles):
            with col: metric(*tl)
        x, y = st.columns(2)
        with x:
            card("Heart & ECG", [p[0] for p in PARAMS if p[7] == "Heart & ECG"], d, extra=f'<div class="row"><span>Rhythm</span><b>{d["rhythm"]}</b></div>' if d.get("rhythm") else f'<div class="row"><span>Rhythm</span><span class="miss">{ml} · enter manually</span></div>')
            card("Cardiac biomarkers", [p[0] for p in PARAMS if p[7] == "Cardiac biomarkers"], d)
            card("Imaging & structure", [p[0] for p in PARAMS if p[7] == "Imaging & structure"], d)
        with y:
            card("Lipid profile", [p[0] for p in PARAMS if p[7] == "Lipid profile"], d)
            card("Metabolic & renal", [p[0] for p in PARAMS if p[7] == "Metabolic & renal"], d)
            chips = "".join(f'<span class="chip">{l}</span>' for k, l in FLAGLAB.items() if d.get(k))
            H(f'<div class="card"><h6>History, findings & symptoms</h6>{chips or "<span class=miss>None documented in the data. Add them on the Patient Profile page.</span>"}</div>')
        if ss.mode == "report" and len(ss.files) > 1:
            st.markdown("**Risk score per uploaded report**")
            pts = [(f["file"], risk({**MISSING, **f["params"]})[0]) for f in ss.files]
            st.line_chart(pd.DataFrame(pts, columns=["report", "score"]).set_index("report"))

def p_waveforms():
    main, d = page("Live Waveforms")
    with main:
        import plotly.graph_objects as go
        if d.get("hr") is None:
            st.info("Heart rate was not found in the reports, so no waveform can be drawn."); return
        st.caption("Note: synthetic ECG/PPG drawn from the extracted heart rate, QTc and ST values — reports contain numbers, not raw signals.")
        hr, qtc, stv = d["hr"], d.get("qtc") or 420, d.get("st_mv") or 0.0
        live = st.toggle("Animate", False)
        def draw():
            t = np.linspace(0, 4, 800) + (time.time() % 4 if live else 0); rr = 60 / hr; ph = (t % rr) / rr
            g = lambda c, w, a: a * np.exp(-((ph - c) ** 2) / (2 * w ** 2)); tc = 0.18 + min(0.5, qtc / 1000 / rr) * 0.8
            ecg = g(.15, .025, .12) - g(.27, .008, .12) + g(.30, .010, 1.1) - g(.33, .010, .25) + stv * (ph > .35) * (ph < tc) + g(tc, .04, .3)
            ppg = g(.25, .07, 1) + g(.5, .06, .35)
            for name, y, col in (("ECG — Lead II", ecg, B), ("PPG — optical", ppg, "#0891b2")):
                f = go.Figure(go.Scatter(x=t, y=y, line=dict(color=col, width=2.5)))
                f.update_layout(title=name, template="plotly_white", height=240, margin=dict(l=10, r=10, t=45, b=10), xaxis_visible=False, yaxis_visible=False)
                st.plotly_chart(f, use_container_width=True)
        st.fragment(run_every=0.5 if live else None)(draw)()
        a, b, c = st.columns(3)
        t, tc = bp_class(d.get("sbp"), d.get("dbp"))
        with a: metric("Blood pressure", f"{fmt(d.get('sbp'))}/{fmt(d.get('dbp'))}", "mmHg", t, tc)
        with b: metric("Physiological stress", fmt(d.get("stress")), "/ 100")
        with c: metric("HRV (RMSSD)", fmt(d.get("hrv")), "ms")

def _set(k, wk): ss.manual[k] = ss[wk]

def prow(label):
    a, b = st.columns([2.4, 3.6], vertical_alignment="center")
    with a: H(f'<div class="plabel">{label}</div>')
    return b

def _seg(k, label, opts):
    wk = "p_" + k; ss[wk] = get(k) if get(k) in opts else None
    if hasattr(st, "segmented_control"): st.segmented_control(label, opts, key=wk, label_visibility="collapsed", on_change=_set, args=(k, wk))
    else: st.radio(label, opts, horizontal=True, index=None, key=wk, label_visibility="collapsed", on_change=_set, args=(k, wk))

def pnum(k, label, unit, lo, hi, step, ph="Enter value"):
    wk = "p_" + k; ss[wk] = None if get(k) is None else float(get(k))
    b = prow(f"{label} ({unit})")
    b.number_input(label, float(lo), float(hi), value=None, step=float(step), format="%g", key=wk, label_visibility="collapsed", placeholder=ph, on_change=_set, args=(k, wk))
    if get(k) is None: b.markdown(f":orange[{miss_label()} · enter manually]")

def pseg(k, label, opts):
    H(f'<div class="plabel" style="margin:14px 0 4px">{label}</div>')
    _seg(k, label, opts)
    if get(k) is None: st.markdown(f":orange[{miss_label()} · select manually]")

def ptoggle(k):
    wk = "p_" + k; ss[wk] = bool(get(k))
    a, b = st.columns([5, 1], vertical_alignment="center")
    with a:
        H(f'<div style="font-weight:600;font-size:.98rem">{FLAGLAB[k]}</div>')
        if get(k) is None and ss.mode in ("report", "scenario"): st.caption(f"{'Not documented in report' if ss.mode == 'report' else 'Not in demo data'}")
    b.toggle(FLAGLAB[k], key=wk, label_visibility="collapsed", on_change=_set, args=(k, wk))
    H('<div style="border-bottom:1px solid #e2e8f0;margin:0 0 6px"></div>')

def p_profile():
    main, d = page("Patient Profile")
    with main:
        st.caption("Patient-reported and clinical inputs. Values found in the uploaded report are filled in; anything missing is marked so you can enter it.")
        with st.expander("Demographics", True):
            b = prow("Name"); ss.p_name = get("name") or ""
            b.text_input("Name", key="p_name", label_visibility="collapsed", placeholder="Patient full name", on_change=_set, args=("name", "p_name"))
            if not get("name"): b.markdown(f":orange[{miss_label()} · enter manually]")
            pnum("age", "Age", "years", 1, 120, 1, "e.g. 52")
            with prow("Sex"): _seg("sex", "Sex", ["Male", "Female"])
            if get("sex") is None: st.markdown(f":orange[{miss_label()} · select manually]")
            pnum("height_cm", "Height", "cm", 100, 230, 1, "e.g. 170"); pnum("weight_kg", "Weight", "kg", 25, 300, 0.5, "e.g. 70")
            bmi = derive(values()).get("bmi")
            b = prow("Body mass index (BMI) (auto)")
            cat = "" if not bmi else "Underweight" if bmi < 18.5 else "Healthy" if bmi < 23 else "Overweight" if bmi < 25 else "Obese range"
            with b:
                H(f'<div style="text-align:right"><span style="font-size:1.6rem;font-weight:800;color:#2563eb">{bmi if bmi else "—"}</span> <span class="mut">kg/m²</span></div>' if bmi
                  else '<div style="text-align:right" class="miss">Not present · enter height and weight</div>')
            H(f'<div class="mut">18.5–22.9 healthy · Regency Healthcare{" · current: " + cat if cat else ""}</div>')
        with st.expander("Habits & lifestyle", True):
            pseg("smoking", "Smoking", ["Never", "Former", "Current"]); pseg("activity", "Physical activity", ["Low", "Moderate", "High"])
            pseg("diet", "Diet quality", ["Low", "Moderate", "High"]); pseg("alcohol", "Alcohol", ["None", "Light", "Moderate", "Heavy"])
        with st.expander("Medical history", True):
            for k in ["diabetes", "htn", "fam_hx", "abd_obesity", "dyslip", "prior_mi", "prior_pci", "prior_cabg", "stroke", "ckd", "afib"]: ptoggle(k)
        with st.expander("ECG and imaging findings"):
            for k in ["lvh", "wma", "lbbb"]: ptoggle(k)
        with st.expander("Symptoms", True):
            for k in ["chest_pain", "dyspnea", "fatigue", "palp", "syncope"]: ptoggle(k)
            pseg("psych_stress", "Psychosocial stress", ["Low", "Moderate", "High"])

def p_lab():
    main, d = page("Lab Report")
    with main:
        st.caption("Values found in your reports. Edit any value to override it; calculated values update automatically.")
        uploader("Add more reports")
        for grp in GROUPS:
            with st.expander(grp, True):
                for col, p in zip(st.columns(3) * 4, [p for p in PARAMS if p[7] == grp]):
                    with col: field(p[0], p[1], p[2], p[3], p[4], p[5])
        rows = [("Non-HDL", "non_hdl", "total − HDL"), ("LDL (Friedewald)", "ldl", "total − HDL − TG/5; unreliable if TG > 400"), ("ApoB (estimated)", "apob", "0.65 × non-HDL + 6.3"),
                ("ApoB / ApoA1", "apob_ratio", "< 0.80 desirable")]
        H('<div class="card"><h6>Calculated lipid values</h6>' + "".join(f'<div class="row"><span>{a}<br><span class="mut">{n}</span></span><b>{fmt(d.get(k))}</b></div>' for a, k, n in rows) + "</div>")

def p_report():
    main, d = page("Patient Report")
    with main:
        if not ss.files: st.info("Upload reports to see the extraction audit."); return
        st.subheader("Extraction audit")
        df = pd.DataFrame([{"file": f["file"], "parameter": h["key"], "value": h["value"], "confidence": h["conf"], "matched text": h["raw"]} for f in ss.files for h in f["hits"]])
        st.dataframe(df, use_container_width=True, hide_index=True)
        if ss.conflicts:
            st.warning("Files disagree on some values; the later upload is used (edit manually to override).")
            st.dataframe(pd.DataFrame(ss.conflicts, columns=["parameter", "earlier", "earlier file", "used", "used file"]), hide_index=True)
        a, b = st.columns(2)
        a.download_button("Download CSV", df.to_csv(index=False), "extracted_parameters.csv")
        b.download_button("Download JSON", json.dumps({k: v for k, v in d.items() if v is not None and not k.startswith("_")}, indent=2, default=str), "patient_parameters.json")

def p_scenarios():
    main, d = page("Scenarios (demo data)", allow_empty=True)
    with main:
        det = detect_scenario(values()) if ss.mode != "empty" else None
        st.caption("The scenario for the current patient is detected automatically from the data. The list below is also available as demo data (Load).")
        for tab, grp in zip(st.tabs(["Healthy", "Coronary artery disease", "Other CVD"]), ["Healthy", "CAD", "Other CVD"]):
            with tab:
                for g, n, desc, hr, s, dia, hrv, _ in (x for x in SCENARIOS if x[0] == grp):
                    t, tc = bp_class(s, dia)
                    H(f'<div class="card"><b>{n}</b> {pill("Detected for this patient", B) if det and det["name"] == n else ""}<div class="mut">{desc}</div><div class="row"><span>HR {hr} bpm · BP {s}/{dia} · HRV {hrv} ms</span>{pill(t, tc)}</div></div>')
                    if st.button("Active" if ss.mode == "scenario" and ss.scn == n else "Load", key=f"l{n}", use_container_width=True): load_scenario(n); st.rerun()

OVERVIEW = st.Page(p_overview, title="Patient Overview", icon=":material/upload_file:", default=True)
DASH = st.Page(p_dashboard, title="Dashboard", icon=":material/dashboard:")
WAVE = st.Page(p_waveforms, title="Live Waveforms", icon=":material/monitor_heart:")
PROFILE = st.Page(p_profile, title="Patient Profile", icon=":material/person:")
LABP = st.Page(p_lab, title="Lab Report", icon=":material/science:")
REPORT = st.Page(p_report, title="Patient Report", icon=":material/description:")
SCEN = st.Page(p_scenarios, title="Scenarios", icon=":material/grid_view:")
PAGES = {"Home": [OVERVIEW], "Monitor": [DASH, WAVE], "Patient": [PROFILE, LABP, REPORT, SCEN]}
try:
    nav, custom = st.navigation(PAGES, position="hidden"), True
except TypeError:  # older Streamlit: use the built-in menu and a larger st.logo
    nav, custom = st.navigation(PAGES), False
if custom:
    with st.sidebar:
        if _LOGO.exists(): st.image(str(_LOGO), width=250)
        for sec, pgs in PAGES.items():
            st.caption(sec.upper())
            for p in pgs: st.page_link(p)
elif _LOGO.exists():
    st.logo(str(_LOGO), size="large"); st.markdown("<style>" + _LOGO_CSS + "</style>", unsafe_allow_html=True)
nav.run()
